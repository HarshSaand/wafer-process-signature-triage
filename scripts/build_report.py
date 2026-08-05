#!/usr/bin/env python3
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'outputs'; FIG=OUT/'figures'; OUT.mkdir(exist_ok=True)
m=json.loads((OUT/'metrics.json').read_text()); audit=json.loads((ROOT/'data/processed/manifest.summary.json').read_text())
NAVY='082B4C'; TEAL='008C95'; GOLD='D6A33A'; PALE='EAF4F4'; GRAY='536777'; WHITE='FFFFFF'

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.82); sec.right_margin=Inches(.82); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
styles=doc.styles
for name,size,color,bold,spaceb,spacea in [('Normal',10.2,'263746',False,0,6),('Title',28,NAVY,True,0,8),('Subtitle',13,GRAY,False,0,12),('Heading 1',17,NAVY,True,14,6),('Heading 2',13,TEAL,True,10,4),('Heading 3',11,GOLD,True,8,3)]:
 s=styles[name]; s.font.name='Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=bold; s.paragraph_format.space_before=Pt(spaceb); s.paragraph_format.space_after=Pt(spacea); s.paragraph_format.line_spacing=1.12
header=sec.header.paragraphs[0]; header.text='WAFER PROCESS-SIGNATURE TRIAGE  |  TECHNICAL REPORT'; header.style=styles['Normal']; header.runs[0].font.size=Pt(8); header.runs[0].font.bold=True; header.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run('Harsh Saand  |  Prototype evidence - not production validation  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for r in footer.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd')) or OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd) if shd.getparent() is None else None
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.style='Table Grid'
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=str(h); shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor.from_string(WHITE); r.font.size=Pt(8.5)
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for i,v in enumerate(row):
   cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ri%2: shade(cells[i],'F3F7F9')
   for r in cells[i].paragraphs[0].runs: r.font.size=Pt(8.3)
 if widths:
  for row in t.rows:
   for c,w in zip(row.cells,widths): c.width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
def bullet(text): doc.add_paragraph(text,style='List Bullet')
def pic(name,width=6.6,caption=None):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(FIG/name),width=Inches(width))
 if caption:
  q=doc.add_paragraph(caption); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.runs[0].italic=True; q.runs[0].font.size=Pt(8.5); q.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
def callout(label,text):
 t=doc.add_table(rows=1,cols=1); t.autofit=False; t.columns[0].width=Inches(6.6); c=t.cell(0,0); shade(c,PALE); p=c.paragraphs[0]; r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL); p.add_run(text); doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Editorial cover pattern.
doc.add_paragraph('A*STAR IME PORTFOLIO PROTOTYPE').runs[0].font.color.rgb=RGBColor.from_string(TEAL)
doc.add_paragraph('Wafer Process-Signature Triage',style='Title')
doc.add_paragraph('Geometry-aware AI for interpretable wafer-map screening and process-excursion review',style='Subtitle')
doc.add_paragraph('HARSH SAAND',style='Heading 2'); doc.add_paragraph('Technical report  |  30 July 2026  |  Locally reproduced on Apple M3 Pro CPU')
pic('system_flow.png',3.85)
callout('Evidence boundary','This is a reproducible classifier and triage prototype evaluated on the shared 902-image Kaggle JPEG derivative. It is not physics simulation, root-cause diagnosis, or validation on A*STAR IME fab data.')
doc.add_page_break()

doc.add_heading('Executive summary',1)
doc.add_paragraph(f"This project demonstrates an end-to-end semiconductor AI workflow: dataset forensics, near-duplicate-aware splitting, a class-balanced baseline, a geometry-aware Cartesian/polar fusion CNN, validation-only calibration, per-class evaluation, spatial signatures, exemplar retrieval, a local interface, tests, and reproducibility scripts. On the untouched 138-image grouped test split, the fusion model achieved macro-F1 {m['macro_f1']:.3f}, balanced accuracy {m['balanced_accuracy']:.3f}, and accuracy {m['accuracy']:.3f}. The engineered-feature logistic baseline achieved macro-F1 {m['baseline_logistic']['macro_f1']:.3f}. These measurements apply only to this derivative and split.")
callout('Core contribution','The model uses two views of each wafer: Cartesian pixels preserve localized/scratch morphology; a differentiable polar unwrap makes radial and annular patterns easier to represent. Engineered spatial proxies and embedding retrieval make outputs more legible to process engineers.')
doc.add_heading('Semiconductor process-analysis problem',1)
doc.add_paragraph('Wafer-sort or metrology maps compress spatial information about a process outcome. Their morphology can prioritize investigation: edge concentration can motivate edge-exclusion or boundary-condition checks; radial structure can motivate thermal, plasma, coating, implant, etch, or CMP-uniformity studies; directional structure can motivate handling or scan-path checks. The same visual signature can have multiple causes, so the prototype ranks pattern evidence and examples rather than asserting root cause.')
doc.add_paragraph('The prototype supports process investigation by quantifying morphology and organizing comparable cases. Process physics, equipment telemetry, recipe context, metrology, and controlled experiments remain necessary to test causality.')
pic('system_flow.png',5.7,'Figure 1. Top-to-bottom implemented prototype flow.')

doc.add_page_break(); doc.add_heading('Dataset provenance and audit',1)
doc.add_paragraph('Source: Kaggle dataset “WM811K Silicon Wafer Map Dataset Image” by muhammedjunayed, downloaded through Kaggle’s official dataset API. The shared archive is a curated image derivative: 902 labeled RGB JPEG heatmaps in nine folders, each 32×32 pixels. It is not the original 811,457-record WM-811K pickle and it does not preserve explicit background/pass/fail integer states, lot names, recipes, tools, process parameters, or train/test metadata. The repository retains the source URL, download script, archive, and provenance notes; users must review Kaggle’s current terms before redistribution.')
table(['Class','Images','Test support'],[(c,audit['classes'][c],m['per_class'][c]['support']) for c in m['per_class']],[2.3,1.5,1.5])
pic('class_counts.png',6.3,'Figure 2. Actual audited class counts in the downloaded archive.')
pic('dataset_examples.png',5.8,'Figure 3. Representative rendered inputs. Brightness/color descriptors are proxies, not exact failed-die counts.')
doc.add_heading('Duplicate and leakage risks',2)
doc.add_paragraph(f"All 902 decoded images had unique SHA-256 hashes. Conservative within-class dHash grouping (Hamming distance <=3) formed {audit['perceptual_groups']} groups. Groups, not individual files, were assigned to partitions, yielding 627 train, 137 validation, and 138 test images. The test partition contains 138 groups; no group crosses partitions. Because the derivative lacks lot IDs and original wafer lineage, lot-level leakage cannot be excluded. This is a candid limitation, not evidence that leakage is absent.")

doc.add_page_break(); doc.add_heading('Methodology',1)
doc.add_heading('Preprocessing and spatial signatures',2)
doc.add_paragraph('JPEGs are decoded as RGB for audit and resized to 64×64 grayscale for modeling. Pixel scaling is fixed to [0,1]. The polar view is generated inside the network with differentiable grid sampling. Proxy descriptors include activated-pixel fraction, centroid, radial mean and spread, core/middle/edge density, quadrant density, anisotropy/orientation, and angular/radial entropy. A saturation/darkness heuristic derives the proxy mask from rendered colors; therefore descriptor names must not be interpreted as exact die-level failure statistics.')
doc.add_heading('Architecture and AI contribution',2)
table(['Component','Implementation','Purpose'],[
 ('Cartesian encoder','3 convolution/BN/SiLU/max-pool stages','Local clusters and scratch morphology'),
 ('Polar encoder','Same compact encoder after polar unwrap','Radial, donut, edge-ring structure'),
 ('Feature branch','LayerNorm + 32-unit MLP','Engineer-readable spatial proxies'),
 ('Fusion','Concatenate; 128-unit MLP; dropout 0.2','Nine-class logits and 128-D retrieval embedding'),
 ('Calibration','Single temperature fit on validation logits','Probability calibration without test tuning')],[1.25,2.65,2.5])
doc.add_paragraph('Training uses class-weighted cross-entropy, AdamW (learning rate 0.002, weight decay 0.0001), batch size 64, seed 42, and 18 epochs. The checkpoint with minimum validation negative log-likelihood was selected; the held-out test split was evaluated once after training. The run used PyTorch 2.13.0 on an Apple M3 Pro CPU. A balanced multinomial logistic regression on engineered proxies provides the sensible fast baseline.')
doc.add_heading('Reproducibility controls',2)
for x in ['Fixed Python, NumPy, and PyTorch seeds.','Near-duplicate groups are indivisible across splits.','Class weighting is derived from the training partition only.','Temperature scaling uses validation logits only.','Retrieval reference embeddings contain training images only.','No augmentation is applied to validation or test inputs.']: bullet(x)

doc.add_page_break(); doc.add_heading('Locally measured results',1)
table(['Model','Accuracy','Balanced accuracy','Macro-F1','NLL','ECE (15 bins)','Brier'],[
 ('Engineered logistic',f"{m['baseline_logistic']['accuracy']:.3f}",f"{m['baseline_logistic']['balanced_accuracy']:.3f}",f"{m['baseline_logistic']['macro_f1']:.3f}",f"{m['baseline_logistic']['negative_log_likelihood']:.3f}",f"{m['baseline_logistic']['ece_15']:.3f}",f"{m['baseline_logistic']['brier_score']:.3f}"),
 ('Fusion CNN',f"{m['accuracy']:.3f}",f"{m['balanced_accuracy']:.3f}",f"{m['macro_f1']:.3f}",f"{m['negative_log_likelihood']:.3f}",f"{m['ece_15']:.3f}",f"{m['brier_score']:.3f}")],[1.45,.75,1.05,.75,.6,.85,.7])
pic('model_comparison.png',5.6,'Figure 4. Test macro-F1 comparison. No confidence interval is claimed for this single split/run.')
doc.add_paragraph('The CNN improves macro-F1 by 0.672 absolute over the engineered baseline. This supports the value of learned spatial morphology on this derivative; it does not isolate the individual contribution of polar versus Cartesian views because a formal ablation was not run. The low baseline also shows that the JPEG proxy-mask descriptors lose important rendered intensity/texture information.')
pic('confusion_matrix.png',6.25,'Figure 5. Held-out confusion matrix (rows=true, columns=predicted).')
pic('per_class_f1.png',6.25,'Figure 6. Per-class F1; the dashed line is macro-F1.')

doc.add_heading('Per-class findings',1)
table(['Class','Precision','Recall','F1','Observed errors'],[(c,f"{v['precision']:.3f}",f"{v['recall']:.3f}",f"{v['f1']:.3f}",('Perfect on this split' if v['f1']==1 else 'See confusion matrix')) for c,v in m['per_class'].items()],[1.25,.75,.75,.75,2.4])
doc.add_paragraph('“none” was perfect on this split and near-full/edge-ring were strongest. Donut, scratch, edge-local, local, and random were harder; notable confusions included donut→random, scratch→edge-local/local, and local→donut. With only 15–16 test examples per class, each error changes recall materially, so production conclusions would be inappropriate.')
doc.add_heading('Calibration and uncertainty',2)
doc.add_paragraph(f"After validation-only temperature scaling, test NLL was {m['negative_log_likelihood']:.3f}, 15-bin ECE {m['ece_15']:.3f}, and multiclass Brier score {m['brier_score']:.3f}. The CLI exposes calibrated probabilities. An operational abstention threshold should be selected against an explicit coverage/risk target on richer validation data; this prototype does not claim fab-grade OOD performance.")

doc.add_page_break(); doc.add_heading('Explainability and process-engineer workflow',1)
doc.add_paragraph('Inference returns the predicted pattern, calibrated class probabilities, entropy, a confidence-based review flag, engineer-readable spatial descriptors, and nearest training images in the learned embedding. These mechanisms answer complementary questions: what class resembles the input, how certain is the model, which geometry is present, and what prior cases look similar? The interface uses cautious wording such as “consistent with” and “investigate.”')
table(['Signature','Possible investigation direction - hypothesis only','Useful corroborating data'],[
 ('Center/donut','Radial thermal, dose, flow, deposition or CMP nonuniformity','Across-wafer thickness/CD, chamber position, recipe step'),
 ('Edge ring/local','Edge exclusion, bevel, clamp, plasma boundary or handling','Edge metrology, bevel inspection, chamber kit state'),
 ('Scratch/directional','Handling, scan path or contact event','Tool logs, robot path, optical review'),
 ('Random/near-full','Stochastic contamination or widespread excursion','Particle data, tool/lot/time correlation, yield history')],[1.25,3.1,2.15])
callout('Interpretability boundary','A saliency map or nearest neighbor can reveal model evidence, but neither establishes a physical mechanism. Root cause requires process context and designed validation.')
doc.add_heading('Application and inference',2)
doc.add_paragraph('The Streamlit dashboard accepts a supported JPEG/PNG or dataset example and presents the triage card. The equivalent CLI is: python scripts/infer.py --image <path> --checkpoint artifacts/model.pt. All processing remains local.')

doc.add_heading('Limitations and validity threats',1)
for x in [
 'Dataset derivative: only 902 rendered JPEGs; headline WM-811K scale must not be attributed to this experiment.',
 'Lost semantics: the original background/pass/fail die states are unavailable, so spatial descriptors are color/intensity proxies.',
 'No fab metadata: no lot grouping, tool, recipe, metrology, time, yield, or process parameters; causal inference is impossible.',
 'Selection and label bias: the derivative’s construction and sampling may make classes easier or less representative than real excursions.',
 'Single run/split: no multi-seed confidence intervals or polar/CNN feature ablations were completed; result uncertainty is unquantified.',
 'OOD: confidence and entropy are triage signals, not validated novelty detection on real IME distributions.',
 'Deployment: production requires fab-specific validation, traceability, monitoring, security, human review, and change control.'
]: bullet(x)

doc.add_page_break(); doc.add_heading('Conclusion',1)
doc.add_paragraph(f"The completed prototype demonstrates a reproducible semiconductor-AI workflow on the shared wafer-map image derivative. The geometry-aware fusion model achieved macro-F1 {m['macro_f1']:.3f} and balanced accuracy {m['balanced_accuracy']:.3f} on the untouched grouped test split, substantially exceeding the engineered-feature baseline. Its calibrated predictions, spatial descriptors, and similar-case retrieval provide an interpretable screening output while preserving a clear boundary between pattern evidence and physical root-cause claims.")
doc.add_paragraph('The strongest evidence is the complete local implementation: official-source download, actual dataset audit, leakage-aware partitioning, baseline and deep model training, held-out evaluation, explainable outputs, CLI and Streamlit inference, automated tests, and traceable artifacts. External validity remains limited by the small rendered JPEG derivative and missing fab metadata.')

doc.add_heading('References and provenance',1)
for x in ['Kaggle shared derivative: https://www.kaggle.com/datasets/muhammedjunayed/wm811k-silicon-wafer-map-dataset-image','Original WM-811K paper: Wu, M.-J. et al., “Wafer Map Failure Pattern Recognition and Similarity Ranking for Large-Scale Data Sets,” IEEE TSM, 2015.','PyTorch documentation: https://pytorch.org/docs/stable/','scikit-learn metrics and logistic regression: https://scikit-learn.org/stable/']: bullet(x)
out=OUT/'wafer_process_signature_triage_report.docx'; doc.save(out); print(out)
